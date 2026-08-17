# Render a project Markdown file as Word. Written 2026-08-16 for
# "Informed Build - Issues for Review", which goes to Gustavo alongside the edit
# docs. The Markdown stays the source of truth; regenerate the .docx from it.
#
#   python3 md2docx.py input.md output.docx
#
# Handles headings, tables, bullets, numbered lists, block quotes, bold, italic and
# inline code, and colours the three status tags (FIX AT SOURCE, LIVE SITE,
# REVIEW). It is deliberately small: if a document needs more than this, the
# document is doing too much.

# One-off: render the issues Markdown as a Word document Neal can send.
# The Markdown at the project root stays the source of truth; regenerate from it.
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = sys.argv[1]; OUT = sys.argv[2]
doc = Document()
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
st.paragraph_format.space_after = Pt(6)

TAGS = {'FIX AT SOURCE': RGBColor(0xB0, 0x00, 0x00),
        'LIVE SITE': RGBColor(0xC0, 0x50, 0x00),
        'REVIEW': RGBColor(0x00, 0x50, 0xA0)}


def runs(par, text):
    """Bold on **...**, colour the three status tags, italics on *...*."""
    for chunk in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)', text):
        if not chunk:
            continue
        if chunk.startswith('**'):
            r = par.add_run(chunk[2:-2]); r.bold = True
            for tag, col in TAGS.items():
                if chunk[2:-2].startswith(tag):
                    r.font.color.rgb = col
        elif chunk.startswith('`'):
            r = par.add_run(chunk[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(9.5)
        elif chunk.startswith('*'):
            r = par.add_run(chunk[1:-1]); r.italic = True
        else:
            par.add_run(chunk)


lines = open(SRC).read().split('\n')
i = 0
while i < len(lines):
    ln = lines[i].rstrip()
    if re.match(r'^\|', ln) and i + 1 < len(lines) and re.match(r'^\|[\s:|-]+\|$', lines[i+1]):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            if not re.match(r'^\|[\s:|-]+\|$', lines[i]):
                rows.append([c.strip() for c in lines[i].strip('|').split('|')])
            i += 1
        t = doc.add_table(rows=0, cols=len(rows[0])); t.style = 'Light Grid Accent 1'
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for ci, val in enumerate(row[:len(cells)]):
                cells[ci].text = ''
                runs(cells[ci].paragraphs[0], val)
                if ri == 0:
                    for r in cells[ci].paragraphs[0].runs:
                        r.bold = True
        doc.add_paragraph()
        continue
    if ln.startswith('# '):
        doc.add_page_break() if doc.paragraphs and len(doc.paragraphs) > 3 else None
        h = doc.add_heading(ln[2:], 1)
    elif ln.startswith('## '):
        doc.add_heading(ln[3:], 2)
    elif ln.startswith('### '):
        doc.add_heading(ln[4:], 3)
    elif ln.startswith('> '):
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.35)
        runs(p, ln[2:]); [setattr(r, 'italic', True) for r in p.runs]
    elif re.match(r'^\s*[-*] ', ln):
        indent = (len(ln) - len(ln.lstrip())) // 2
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.left_indent = Inches(0.3 + 0.25 * indent)
        runs(p, re.sub(r'^\s*[-*] ', '', ln))
    elif re.match(r'^\s*\d+\. ', ln):
        p = doc.add_paragraph(style='List Number')
        runs(p, re.sub(r'^\s*\d+\. ', '', ln))
    elif ln.strip() == '---':
        pass
    elif ln.strip():
        buf = [ln]
        while i + 1 < len(lines) and lines[i+1].strip() and not re.match(
                r'^(#|\||>|\s*[-*] |\s*\d+\. |---)', lines[i+1]):
            i += 1; buf.append(lines[i].strip())
        runs(doc.add_paragraph(), ' '.join(buf))
    i += 1

doc.save(OUT)
print('wrote', OUT)
